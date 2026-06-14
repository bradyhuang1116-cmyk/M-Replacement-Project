"""图纸处理日志记录表 — 开放给 PLM 系统后台访问。

字段对齐客户文档《光栅图处理软件与PLM系统的交互机制》中的 R_V_TD_FILEPATH 表。
处理方式标记 OCR 字段：
  - O = 使用了 OCR 识别（扫描件路径，method='ocr'）
  - N = 未使用 OCR（矢量 PDF 直接处理，method='vector'）

PLM 来源字段（docnumber/work_seq）单机处理时拿不到，留空；
对方 PLM 对接层可在回写客户 Oracle 时补齐。
"""
from __future__ import annotations

import os
import sqlite3
import threading
from datetime import datetime
from pathlib import Path

from config import PROCESS_LOG_DB_PATH

# ── 测试期开关 ──────────────────────────────────────────────
# 处理日志除写入 SQLite 表外，同时追加一份 CSV 文件（data/process_log.csv），
# 便于测试期直接打开查看 O/N 等记录。
# 【测试完毕改回】：把下面 _WRITE_CSV_LOG 设为 False（或删除 _append_csv 调用），
# 生产以 SQLite 表 process_log 为准、供 PLM 读取，无需 CSV。
_WRITE_CSV_LOG = False
_CSV_LOG_PATH = os.path.join(os.path.dirname(PROCESS_LOG_DB_PATH), "process_log.csv")

_SCHEMA = """
CREATE TABLE IF NOT EXISTS process_log (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    docnumber     TEXT,            -- 文件ID (R_V_TD_FILEPATH.DOCNUMBER)，PLM 来源，可空
    drawing_no    TEXT,            -- 图号 (TH)
    revision      TEXT,            -- 版本 (BBH)
    filename      TEXT,            -- 文件名 (FILENAME)
    ocr_flag      TEXT NOT NULL,   -- 处理方式：O=用OCR / N=未用OCR
    work_seq      TEXT,            -- 接收批次号 (WORK_SEQ)，PLM 来源，可空
    status        TEXT NOT NULL,   -- success / failed
    process_date  TEXT NOT NULL    -- 处理日期时间
);

CREATE INDEX IF NOT EXISTS idx_pl_drawing ON process_log(drawing_no);
CREATE INDEX IF NOT EXISTS idx_pl_date ON process_log(process_date);
"""


def method_to_ocr_flag(method: str) -> str:
    """处理方式 → OCR 标记。method='ocr'→O（用了OCR）；'vector'→N（未用OCR）。"""
    return "O" if method == "ocr" else "N"


class ProcessLog:
    """SQLite 图纸处理日志表。线程安全（写操作加锁）。"""

    def __init__(self, db_path: str | Path | None = None):
        self._db_path = str(db_path or PROCESS_LOG_DB_PATH)
        Path(self._db_path).parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self._bootstrap()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(self._db_path, timeout=30, isolation_level=None)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode = WAL")
        return conn

    def _bootstrap(self):
        with self._connect() as conn:
            conn.executescript(_SCHEMA)

    def record(
        self,
        drawing_no: str | None,
        revision: str | None,
        ocr_flag: str,
        status: str = "success",
        filename: str | None = None,
        docnumber: str | None = None,
        work_seq: str | None = None,
    ) -> int:
        """写入一条处理日志。返回行 id。"""
        with self._lock, self._connect() as conn:
            cur = conn.execute(
                """
                INSERT INTO process_log
                    (docnumber, drawing_no, revision, filename,
                     ocr_flag, work_seq, status, process_date)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (docnumber, drawing_no, revision, filename,
                 ocr_flag, work_seq, status,
                 datetime.now().isoformat(timespec="seconds")),
            )
            row_id = cur.lastrowid
        if _WRITE_CSV_LOG:
            self._append_csv(drawing_no, revision, filename, ocr_flag,
                             work_seq, docnumber, status)
        return row_id

    def _append_csv(self, drawing_no, revision, filename, ocr_flag,
                    work_seq, docnumber, status):
        """测试期：追加一行到 CSV，便于直接查看。生产可关闭（见文件头 _WRITE_CSV_LOG）。"""
        try:
            new = not os.path.exists(_CSV_LOG_PATH)
            os.makedirs(os.path.dirname(_CSV_LOG_PATH), exist_ok=True)
            with open(_CSV_LOG_PATH, "a", encoding="utf-8-sig", newline="") as f:
                import csv
                w = csv.writer(f)
                if new:
                    w.writerow(["处理日期", "图号", "版本", "文件名",
                                "处理方式(O/N)", "状态", "文件ID", "批次号"])
                w.writerow([datetime.now().isoformat(timespec="seconds"),
                            drawing_no or "", revision or "", filename or "",
                            ocr_flag, status, docnumber or "", work_seq or ""])
        except Exception:
            pass

    def list_recent(self, limit: int = 100) -> list[dict]:
        items, _ = self.query(limit=limit, offset=0)
        return items

    def query(
        self,
        drawing_no: str | None = None,
        revision: str | None = None,
        mode: str | None = None,
        status: str | None = None,
        date_from: str | None = None,
        date_to: str | None = None,
        limit: int = 100,
        offset: int = 0,
    ) -> tuple[list[dict], int]:
        """分页查询处理日志。mode 对应 ocr_flag（O/N）。"""
        clauses: list[str] = []
        params: list[object] = []

        if drawing_no:
            clauses.append("drawing_no LIKE ?")
            params.append(f"%{drawing_no}%")
        if revision:
            clauses.append("revision LIKE ?")
            params.append(f"%{revision}%")
        if mode:
            clauses.append("ocr_flag = ?")
            params.append(mode.upper())
        if status:
            clauses.append("status = ?")
            params.append(status)
        if date_from:
            clauses.append("process_date >= ?")
            params.append(date_from)
        if date_to:
            clauses.append("process_date <= ?")
            params.append(f"{date_to}T23:59:59")

        where = f"WHERE {' AND '.join(clauses)}" if clauses else ""
        with self._connect() as conn:
            total = conn.execute(
                f"SELECT COUNT(*) FROM process_log {where}", params
            ).fetchone()[0]
            rows = conn.execute(
                f"""
                SELECT * FROM process_log
                {where}
                ORDER BY id DESC
                LIMIT ? OFFSET ?
                """,
                [*params, limit, offset],
            ).fetchall()
        return [dict(r) for r in rows], int(total)

    @staticmethod
    def to_xlsx(rows: list[dict]) -> bytes:
        """将日志行导出为带样式的 XLSX 字节流。

        样式与前端暗色仪表盘主题对齐。
        """
        import io

        try:
            from openpyxl import Workbook
            from openpyxl.styles import (
                Font, PatternFill, Alignment, Border, Side,
            )
        except ImportError:
            raise RuntimeError(
                "openpyxl is not installed. Run: pip install openpyxl"
            )

        wb = Workbook()
        ws = wb.active
        ws.title = "Process Logs"

        # ── 样式常量（匹配前端暗色主题） ──
        HEADER_FILL = PatternFill(start_color="262626", end_color="262626", fill_type="solid")
        HEADER_FONT = Font(name="Inter", bold=True, color="F5F5F5", size=11)
        HEADER_ALIGNMENT = Alignment(horizontal="left", vertical="center")

        ROW_FILL_EVEN = PatternFill(start_color="171717", end_color="171717", fill_type="solid")
        ROW_FILL_ODD = PatternFill(start_color="1C1C1C", end_color="1C1C1C", fill_type="solid")
        ROW_FONT = Font(name="Inter", color="E5E5E5", size=10)

        SUCCESS_FILL = PatternFill(start_color="0A2E1A", end_color="0A2E1A", fill_type="solid")
        SUCCESS_FONT = Font(name="Inter", color="6EE7B7", bold=True, size=10)

        FAILED_FILL = PatternFill(start_color="2E0A0A", end_color="2E0A0A", fill_type="solid")
        FAILED_FONT = Font(name="Inter", color="FDA5A5", bold=True, size=10)

        MODE_OCR_FILL = PatternFill(start_color="1A1A3E", end_color="1A1A3E", fill_type="solid")
        MODE_OCR_FONT = Font(name="Inter", color="A5B4FC", size=10)
        MODE_PDF_FILL = PatternFill(start_color="1A2E1A", end_color="1A2E1A", fill_type="solid")
        MODE_PDF_FONT = Font(name="Inter", color="6EE7B7", size=10)

        THIN_BORDER = Border(
            left=Side(style="thin", color="333333"),
            right=Side(style="thin", color="333333"),
            top=Side(style="thin", color="333333"),
            bottom=Side(style="thin", color="333333"),
        )

        # ── 列配置 ──
        columns = [
            ("Drawing", 42),
            ("Status", 16),
            ("Mode", 18),
            ("Revision", 12),
            ("Processed At", 22),
        ]

        # ── 写入表头 ──
        for col_idx, (header, width) in enumerate(columns, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGNMENT
            cell.border = THIN_BORDER
            ws.column_dimensions[chr(64 + col_idx)].width = width

        ws.row_dimensions[1].height = 32

        # ── 写入数据行 ──
        for i, row in enumerate(rows):
            excel_row = i + 2
            drawing_no = (row.get("drawing_no") or "").strip()
            filename = (row.get("filename") or "").strip()
            drawing = f"{drawing_no}  ({filename})" if drawing_no and filename else (drawing_no or filename or "—")

            raw_status = (row.get("status") or "").strip()
            status_label = "Success" if raw_status == "success" else "Failed"

            ocr = (row.get("ocr_flag") or "").strip()
            mode_label = "VLM OCR" if ocr == "O" else "Vector PDF"

            revision = row.get("revision") or "—"

            raw_date = row.get("process_date") or ""
            if raw_date:
                try:
                    dt = datetime.fromisoformat(raw_date)
                    formatted_date = dt.strftime("%Y-%m-%d  %H:%M:%S")
                except ValueError:
                    formatted_date = raw_date
            else:
                formatted_date = "—"

            row_fill = ROW_FILL_EVEN if i % 2 == 0 else ROW_FILL_ODD

            # ── Drawing ──
            c = ws.cell(row=excel_row, column=1, value=drawing)
            c.font = ROW_FONT
            c.fill = row_fill
            c.alignment = Alignment(vertical="center")
            c.border = THIN_BORDER

            # ── Status ──
            c = ws.cell(row=excel_row, column=2, value=status_label)
            if raw_status == "success":
                c.font = SUCCESS_FONT
                c.fill = SUCCESS_FILL
            else:
                c.font = FAILED_FONT
                c.fill = FAILED_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = THIN_BORDER

            # ── Mode ──
            c = ws.cell(row=excel_row, column=3, value=mode_label)
            if ocr == "O":
                c.font = MODE_OCR_FONT
                c.fill = MODE_OCR_FILL
            else:
                c.font = MODE_PDF_FONT
                c.fill = MODE_PDF_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = THIN_BORDER

            # ── Revision ──
            c = ws.cell(row=excel_row, column=4, value=revision)
            c.font = ROW_FONT
            c.fill = row_fill
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = THIN_BORDER

            # ── Processed At ──
            c = ws.cell(row=excel_row, column=5, value=formatted_date)
            c.font = Font(name="Inter", color="A3A3A3", size=10)
            c.fill = row_fill
            c.alignment = Alignment(vertical="center")
            c.border = THIN_BORDER

            ws.row_dimensions[excel_row].height = 28

        # ── 冻结首行 ──
        ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read()
